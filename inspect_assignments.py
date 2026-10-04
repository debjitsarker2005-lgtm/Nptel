import os
import sys
import docx
from pypdf import PdfReader

# Force UTF-8 output on Windows terminal
sys.stdout.reconfigure(encoding='utf-8')

def inspect_files():
    print("--- Inspecting docx and pdf files ---")
    for fname in sorted(os.listdir("assignments")):
        path = os.path.join("assignments", fname)
        if fname.endswith(".docx"):
            doc = docx.Document(path)
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            print(f"\n==============================\n[DOCX] {fname}: {len(doc.paragraphs)} paragraphs, {len(full_text)} chars")
            print("Preview:\n", full_text[:500])
        elif fname.endswith(".pdf"):
            reader = PdfReader(path)
            text = ""
            for p in reader.pages:
                text += (p.extract_text() or "") + "\n--- PAGE BREAK ---\n"
            print(f"\n==============================\n[PDF] {fname}: {len(reader.pages)} pages, {len(text)} chars")
            print("Preview:\n", text[:500].strip())

if __name__ == "__main__":
    inspect_files()
